# -*- coding: utf-8 -*-
"""One-shot patch: align Docs/用户手册.docx with the reverted (3a5fc3a) SRAM layout.

The a956029 memory refactor was reverted; this script rewrites the manual's
task-stack / video-pool wording, fixes the memory-map table, and swaps the
fig_mem / fig_video images with the regenerated PNGs.
"""
import os
import zipfile

import docx

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
DOCX = os.path.join(ROOT, "Docs", "用户手册.docx")
DIAG = os.path.join(ROOT, "Docs", "images", "diagrams")

REPL = [
    # task stacks come from CCM again
    ("线程栈从 SRAMSTACK 通过 sys_thread_stack_alloc() 获取",
     "线程栈从 CCM（SRAMCCM）通过 sys_thread_stack_alloc() 获取"),
    ("栈通过 sys_thread_stack_alloc() 取自 SRAMSTACK：",
     "栈通过 sys_thread_stack_alloc() 取自 CCM（SRAMCCM）："),
    ("/* 栈取自 SRAMSTACK */", "/* 栈取自 CCM */"),
    ("三个内存池由统一分配器管理", "两个内存池由统一分配器管理"),
    # NES stripe is a static main-SRAM array again
    ("模拟器使用 SRAMEX 对齐条带缓冲向屏幕输出；退出时回收运行时缓冲并恢复 GUI",
     "模拟器使用主 SRAM 静态条带缓冲向屏幕输出；退出时释放运行时缓冲并恢复 GUI"),
    ("渲染输出使用 16 行 × 2 的 LCD 双 stripe，缓冲从 SRAMEX 以 32 字节对齐分配，退出时回收。",
     "渲染输出使用 16 行 × 2 的 LCD 双 stripe，缓冲为主 SRAM 中的静态数组。"),
    # image viewer JPEG path (default JMALLOC -> SRAMEX)
    ("LibJPEG 解码上下文和 arena 使用 CPU-only 的 CCM，显示条带和 DMA 缓冲使用可 DMA 的 SRAMEX/片上视频池。",
     "LibJPEG 解码上下文和 arena 走默认 SRAMEX 堆分配，显示条带和 DMA 缓冲同样使用可 DMA 的 SRAMEX。"),
    # video pipeline wording
    ("Video_MJPEG.c 使用 LibJPEG 流式输入和 CCM arena",
     "Video_MJPEG.c 使用 LibJPEG 流式输入和片上视频池 arena"),
    ("与 JPEG 解码在 CCM 中的计算重叠执行",
     "与 JPEG 解码在片上 SRAM 中的计算重叠执行"),
    ("CCM arena 与片上视频池的峰值和申请失败",
     "视频池 arena 峰值与视频池分配失败计数"),
    # sysstats description
    ("CPU、任务栈和各内存池使用率", "CPU、CCM/SRAMEX 堆使用率"),
]


def iter_all_paragraphs(document):
    for p in document.paragraphs:
        yield p
    for t in document.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p


def replace_in_para(p, old, new):
    runs = p.runs
    texts = [r.text for r in runs]
    full = "".join(texts)
    idx = full.find(old)
    if idx < 0:
        return False
    end = idx + len(old)
    pos = 0
    replaced_first = False
    for r, t in zip(runs, texts):
        s, e = pos, pos + len(t)
        pos = e
        if e <= idx or s >= end:
            continue
        head = t[:max(0, idx - s)] if s < idx else ""
        tail = t[max(0, end - s):] if e > end else ""
        if not replaced_first:
            r.text = head + new + tail
            replaced_first = True
        else:
            r.text = tail
    return True


def set_cell_text(cell, text):
    paras = cell.paragraphs
    p = paras[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)
    for extra in paras[1:]:
        for r in extra.runs:
            r.text = ""


def main():
    doc = docx.Document(DOCX)

    counts = {old: 0 for old, _ in REPL}
    for p in iter_all_paragraphs(doc):
        for old, new in REPL:
            while replace_in_para(p, old, new):
                counts[old] += 1

    for old, _ in REPL:
        print(("OK  " if counts[old] else "MISS") + " x%d  %s" % (counts[old], old[:38]))

    # ---- memory map table: fix cells, drop the SRAMSTACK row ----
    target = None
    for t in doc.tables:
        if t.rows and t.rows[0].cells[0].text.strip() == "区域":
            texts = " ".join(c.text for row in t.rows for c in row.cells)
            if "SRAMSTACK" in texts:
                target = t
                break
    print("memory table:", "found" if target is not None else "MISS")
    if target is not None:
        for row in list(target.rows):
            name = row.cells[0].text.strip()
            if name == "SRAMSTACK":
                row._tr.getparent().remove(row._tr)
            elif name == "普通主 SRAM":
                set_cell_text(row.cells[1], "0x20000000..0x2000BFFF / 48 KB")
                set_cell_text(row.cells[2], ".data/.bss、MSP、C heap、FreeRTOS 控制块、NES LCD stripe 静态缓冲。")
            elif name == "片上 Video_Memory":
                set_cell_text(row.cells[1], "0x2000C000..0x2001FFFF / 80 KB")
                set_cell_text(row.cells[2], "VideoTask 栈、FIL、AVI 上下文、JPEG source、LibJPEG 上下文 + 38 KB arena、RGB 行缓冲、LCD DMA 条带。")
            elif name == "SRAMCCM":
                set_cell_text(row.cells[1], "0x10000000 起 / 64 KB")
                set_cell_text(row.cells[2], "RTOS 任务栈（defaultTask 等）、CPU-only 热数据；严禁 DMA。")
            elif name == "SRAMEX":
                set_cell_text(row.cells[2], "LVGL 堆/帧缓冲、PCM FIFO、I2S DMA、GUI 上下文。")

    doc.save(DOCX)

    # ---- swap the two regenerated figures inside the zip ----
    media = {
        "word/media/image12.png": os.path.join(DIAG, "fig_mem.png"),
        "word/media/image22.png": os.path.join(DIAG, "fig_video.png"),
    }
    tmp = DOCX + ".tmp"
    with zipfile.ZipFile(DOCX) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in media:
                with open(media[item.filename], "rb") as f:
                    data = f.read()
                print("swapped", item.filename, "<-", os.path.basename(media[item.filename]))
            zout.writestr(item, data)
    os.replace(tmp, DOCX)
    print("done:", DOCX)


if __name__ == "__main__":
    main()
