# -*- coding: utf-8 -*-
"""Round-2 repairs after first judge pass:
1. rebuild TOC region with a clean, correctly-placed TOC field
2. LVGL period 10 ms -> 1 ms (matches osDelay(1) in firmware)
3. keepNext on every image paragraph so image+caption stay together
4. hexdump range typo 1..65536 -> 1~65536
5. quote the SHELL_EXPORT_CMD desc string in the code example
6. 17.2 table wording: 已解码 -> 送显
7. swap embedded PNG blobs for the regenerated (larger-font) diagrams
"""
import copy
import os
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

ROOT = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
SRC = os.path.join(ROOT, "Docs", "用户手册.docx")
DIAG = os.path.join(ROOT, "Docs", "images", "diagrams")
doc = docx.Document(SRC)

# ---------------------------------------------------------------- 1. TOC
RPR_BODY = ('<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:eastAsia="Microsoft YaHei"/>'
            '<w:sz w:val="22"/><w:lang w:val="zh-CN" w:eastAsia="zh-CN"/>')
RPR_HINT = ('<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri" w:eastAsia="Microsoft YaHei"/>'
            '<w:i/><w:color w:val="8A94A6"/><w:sz w:val="19"/>'
            '<w:lang w:val="zh-CN" w:eastAsia="zh-CN"/>')

paras = doc.paragraphs
toc_head = next(p for p in paras if p.style.name == "Heading 1" and p.text.strip() == "目录")
part1 = next(p for p in paras if p.style.name == "Heading 1" and p.text.strip().startswith("第一部分"))

seen, doomed = False, []
for p in paras:
    if p._p is toc_head._p:
        seen = True
        continue
    if p._p is part1._p:
        break
    if seen:
        doomed.append(p)
for p in doomed:
    p._p.getparent().remove(p._p)

intro = parse_xml(
    '<w:p %s><w:pPr><w:spacing w:before="0" w:after="120" w:line="300" w:lineRule="auto"/></w:pPr>'
    '<w:r><w:rPr>%s</w:rPr><w:t>本手册按用户操作、二次开发和系统架构三部分组织。目录由标题样式自动生成，'
    '也可使用 Word 导航窗格（视图 → 导航窗格）按标题层级快速跳转。</w:t></w:r></w:p>'
    % (nsdecls('w'), RPR_BODY))
toc_field = parse_xml(
    '<w:p %s><w:pPr><w:spacing w:before="120" w:after="120" w:line="300" w:lineRule="auto"/></w:pPr>'
    '<w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
    '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h \\z \\u </w:instrText></w:r>'
    '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
    '<w:r><w:rPr><w:i/><w:color w:val="8A94A6"/><w:sz w:val="21"/></w:rPr>'
    '<w:t>（目录将在打开文档时自动生成；若未显示页码，请右键此处选择“更新域”）</w:t></w:r>'
    '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>' % nsdecls('w'))
hint = parse_xml(
    '<w:p %s><w:pPr><w:spacing w:before="0" w:after="200" w:line="240" w:lineRule="auto"/></w:pPr>'
    '<w:r><w:rPr>%s</w:rPr><w:t>提示：页码以实际排版为准，更新域后自动刷新。</w:t></w:r></w:p>'
    % (nsdecls('w'), RPR_HINT))
toc_head._p.addnext(intro)
intro.addnext(toc_field)
toc_field.addnext(hint)
print("1. TOC rebuilt, removed", len(doomed), "paragraphs")

# ------------------------------------------------------- 2/4/5/6 text fixes
def fix_runs(container_paras, pairs):
    n = 0
    for p in container_paras:
        for r in p.runs:
            for old, new in pairs:
                if old in r.text:
                    r.text = r.text.replace(old, new)
                    n += 1
    return n

n2 = fix_runs(doc.paragraphs, [("每 10 ms 调用 lv_timer_handler", "每 1 ms 调用 lv_timer_handler")])
print("2. LVGL period fixed:", n2)

cell_fix = 0
for t in doc.tables:
    for row in t.rows:
        for cell in row.cells:
            cell_fix += fix_runs(cell.paragraphs, [
                ("1..65536", "1~65536"),
                ("已解码 2407", "送显 2407"),
                ("已解码 1226", "送显 1226"),
            ])
print("4/6. table text fixed:", cell_fix)

code_fix = fix_runs(doc.paragraphs, [('peek internal SRAM)', '"peek internal SRAM"')])
print("5. code string quoted:", code_fix)

# ------------------------------------------------------------- 3. keepNext
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
kept = 0
for p in doc.paragraphs:
    if p._p.find('.//' + qn('w:drawing')) is not None:
        ppr = p._p.get_or_add_pPr()
        if ppr.find(qn('w:keepNext')) is None:
            kn = ppr.makeelement(qn('w:keepNext'), {})
            style_el = ppr.find(qn('w:pStyle'))
            if style_el is not None:
                style_el.addnext(kn)
            else:
                ppr.insert(0, kn)
            kept += 1
print("3. keepNext added to", kept, "image paragraphs")

# --------------------------------------------------- 7. swap diagram blobs
shapes = list(doc.inline_shapes)
assert len(shapes) == 23, len(shapes)
# document figure order; None = original photo (untouched)
order = [None, None, "fig_clock.png", None, "fig_nes.png", None, "fig_music.png",
         None, None, None, "fig_font.png", "fig_mem.png", None, "fig_boot.png",
         "fig_storage.png", "fig_sdio.png", "fig_page.png", "fig_handshake.png",
         None, "fig_iap.png", "fig_axf.png", None, "fig_video.png"]
from PIL import Image
import io
for idx, (shape, fname) in enumerate(zip(shapes, order)):
    if fname is None:
        continue
    blip = shape._inline.graphic.graphicData.pic.blipFill.blip
    part = doc.part.related_parts[blip.embed]
    data = open(os.path.join(DIAG, fname), "rb").read()
    with Image.open(io.BytesIO(data)) as im:
        nw, nh = im.size
    ratio_doc = shape.width / shape.height
    ratio_png = nw / nh
    assert abs(ratio_doc - ratio_png) / ratio_png < 0.01, (fname, ratio_doc, ratio_png)
    part._blob = data
print("7. blobs swapped for 13 diagrams")

doc.save(SRC)
print("saved")
